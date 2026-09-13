from docx import Document

document = Document()
document.add_heading('Test DOCX File for Verification', 0)
document.add_paragraph('This is a test document to verify mammoth.js parsing functionality in NoteDigest.')
document.add_paragraph('It contains some dummy text that should be extracted and summarized.')
document.save('sample.docx')
print("Created sample.docx")
